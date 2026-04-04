"""Developer Mode attachment parsing helpers."""

from __future__ import annotations

import csv
import shutil
import subprocess
from io import StringIO
from pathlib import Path

from PIL import Image
from PyPDF2 import PdfReader

try:
    import openpyxl
except ImportError:  # pragma: no cover - optional dependency
    openpyxl = None

try:
    import xlrd
except ImportError:  # pragma: no cover - optional dependency
    xlrd = None


TEXT_FILE_SUFFIXES = {
    ".c",
    ".cc",
    ".cpp",
    ".cs",
    ".css",
    ".dart",
    ".env",
    ".go",
    ".h",
    ".hpp",
    ".html",
    ".ini",
    ".java",
    ".js",
    ".json",
    ".jsx",
    ".kt",
    ".kts",
    ".lua",
    ".md",
    ".php",
    ".py",
    ".rb",
    ".rs",
    ".sh",
    ".sql",
    ".svg",
    ".swift",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
    ".zsh",
}
TEXT_MIME_PREFIXES = ("text/",)
TEXT_MIME_TYPES = {
    "application/javascript",
    "application/json",
    "application/sql",
    "application/toml",
    "application/typescript",
    "application/x-httpd-php",
    "application/x-python-code",
    "application/xml",
}
IMAGE_FILE_SUFFIXES = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
CSV_FILE_SUFFIXES = {".csv", ".tsv"}
PDF_FILE_SUFFIXES = {".pdf"}
EXCEL_FILE_SUFFIXES = {".xls", ".xlsx", ".xlsm", ".xltm", ".xltx"}
CSV_MIME_TYPES = {
    "application/csv",
    "application/vnd.ms-excel",
    "text/csv",
    "text/tab-separated-values",
}
EXCEL_MIME_TYPES = {
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.template",
    "application/vnd.ms-excel.template.macroenabled.12",
}
PDF_MIME_TYPES = {"application/pdf"}
MAX_TABLE_ROWS = 200
MAX_TABLE_COLUMNS = 20
MAX_WORKBOOK_SHEETS = 6
MAX_PDF_PAGES = 20
TESSERACT_TIMEOUT_SECONDS = 30


def developer_input_support_text() -> str:
    """Return a short user-facing description of supported developer attachments."""
    return (
        "فایل‌های متنی/کد، عکس و اسکرین‌شات، PDF، CSV و فایل‌های Excel "
        "(.xls/.xlsx/.xlsm/.xltx/.xltm)"
    )


def is_supported_developer_input_file(file_name: str, mime_type: str | None) -> bool:
    """Return whether a document looks like a supported Developer Mode attachment."""
    suffix = Path(file_name).suffix.lower()
    if suffix in TEXT_FILE_SUFFIXES | IMAGE_FILE_SUFFIXES | CSV_FILE_SUFFIXES | PDF_FILE_SUFFIXES | EXCEL_FILE_SUFFIXES:
        return True
    if not mime_type:
        return False
    lowered = mime_type.lower()
    if lowered in TEXT_MIME_TYPES | CSV_MIME_TYPES | PDF_MIME_TYPES | EXCEL_MIME_TYPES:
        return True
    if any(lowered.startswith(prefix) for prefix in TEXT_MIME_PREFIXES):
        return True
    return lowered.startswith("image/")


def read_developer_input_attachment(
    path: Path,
    *,
    file_name: str,
    mime_type: str | None = None,
    max_chars: int,
) -> tuple[str, bool]:
    """Read one Developer Mode attachment and normalize it into text."""
    suffix = Path(file_name).suffix.lower()
    if suffix in IMAGE_FILE_SUFFIXES or (mime_type or "").lower().startswith("image/"):
        text = _read_image_attachment(path, file_name=file_name)
    elif suffix in PDF_FILE_SUFFIXES or (mime_type or "").lower() in PDF_MIME_TYPES:
        text = _read_pdf_attachment(path, file_name=file_name)
    elif suffix in CSV_FILE_SUFFIXES or (mime_type or "").lower() in CSV_MIME_TYPES:
        text = _read_csv_attachment(path, file_name=file_name)
    elif suffix in EXCEL_FILE_SUFFIXES or (mime_type or "").lower() in EXCEL_MIME_TYPES:
        text = _read_excel_attachment(path, file_name=file_name)
    else:
        text = _read_text_attachment(path, file_name=file_name)

    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        normalized = f"[Attachment: {Path(file_name).name}]\nNo readable text could be extracted."
    truncated = len(normalized) > max_chars
    if truncated:
        normalized = normalized[:max_chars].rstrip() + "\n...[truncated]"
    return normalized, truncated


def _read_text_attachment(path: Path, *, file_name: str) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw:
        raise RuntimeError(
            f"{Path(file_name).name} looks binary and cannot be analyzed as plain source text."
        )
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("utf-8", errors="replace")
    return f"[Text attachment: {Path(file_name).name}]\n{text}"


def _read_pdf_attachment(path: Path, *, file_name: str) -> str:
    reader = PdfReader(str(path))
    total_pages = len(reader.pages)
    lines = [f"[PDF attachment: {Path(file_name).name}]", f"Pages: {total_pages}"]
    extracted_any = False
    for page_index, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
        page_text = (page.extract_text() or "").strip()
        if not page_text:
            continue
        extracted_any = True
        lines.append("")
        lines.append(f"Page {page_index}:")
        lines.append(page_text)
    if total_pages > MAX_PDF_PAGES:
        lines.append("")
        lines.append(f"... {total_pages - MAX_PDF_PAGES} more pages omitted.")
    if not extracted_any:
        lines.append("")
        lines.append("No machine-readable text was extracted from this PDF.")
    return "\n".join(lines)


def _read_csv_attachment(path: Path, *, file_name: str) -> str:
    text = _decode_tabular_bytes(path.read_bytes())
    sample = text[:4096]
    delimiter = "," if path.suffix.lower() != ".tsv" else "\t"
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    lines = [f"[CSV attachment: {Path(file_name).name}]"]
    row_count = 0
    truncated_rows = False
    for row_count, row in enumerate(reader, start=1):
        if row_count > MAX_TABLE_ROWS:
            truncated_rows = True
            break
        normalized = [_stringify_cell(cell) for cell in row[:MAX_TABLE_COLUMNS]]
        lines.append(" | ".join(normalized) if normalized else "(empty row)")
    if truncated_rows:
        lines.append(f"... rows after {MAX_TABLE_ROWS} were omitted.")
    return "\n".join(lines)


def _read_excel_attachment(path: Path, *, file_name: str) -> str:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return _read_legacy_excel_attachment(path, file_name=file_name)
    if openpyxl is None:
        raise RuntimeError("openpyxl is required to read .xlsx/.xlsm Excel files.")

    workbook = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    lines = [f"[Excel attachment: {Path(file_name).name}]"]
    sheet_names = workbook.sheetnames[:MAX_WORKBOOK_SHEETS]
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        lines.append("")
        lines.append(f"Sheet: {sheet_name}")
        row_counter = 0
        for row_counter, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_counter > MAX_TABLE_ROWS:
                lines.append(f"... rows after {MAX_TABLE_ROWS} were omitted.")
                break
            normalized = [_stringify_cell(cell) for cell in row[:MAX_TABLE_COLUMNS]]
            lines.append(" | ".join(normalized) if normalized else "(empty row)")
    if len(workbook.sheetnames) > MAX_WORKBOOK_SHEETS:
        lines.append("")
        lines.append(f"... {len(workbook.sheetnames) - MAX_WORKBOOK_SHEETS} more sheets omitted.")
    return "\n".join(lines)


def _read_legacy_excel_attachment(path: Path, *, file_name: str) -> str:
    if xlrd is None:
        raise RuntimeError("xlrd is required to read .xls Excel files.")
    workbook = xlrd.open_workbook(str(path))
    lines = [f"[Excel attachment: {Path(file_name).name}]"]
    for sheet_index in range(min(workbook.nsheets, MAX_WORKBOOK_SHEETS)):
        sheet = workbook.sheet_by_index(sheet_index)
        lines.append("")
        lines.append(f"Sheet: {sheet.name}")
        for row_index in range(min(sheet.nrows, MAX_TABLE_ROWS)):
            row = sheet.row_values(row_index)[:MAX_TABLE_COLUMNS]
            normalized = [_stringify_cell(cell) for cell in row]
            lines.append(" | ".join(normalized) if normalized else "(empty row)")
        if sheet.nrows > MAX_TABLE_ROWS:
            lines.append(f"... rows after {MAX_TABLE_ROWS} were omitted.")
    if workbook.nsheets > MAX_WORKBOOK_SHEETS:
        lines.append("")
        lines.append(f"... {workbook.nsheets - MAX_WORKBOOK_SHEETS} more sheets omitted.")
    return "\n".join(lines)


def _read_image_attachment(path: Path, *, file_name: str) -> str:
    with Image.open(path) as image:
        width, height = image.size
        image_format = image.format or path.suffix.lstrip(".").upper() or "unknown"

    lines = [
        f"[Image attachment: {Path(file_name).name}]",
        f"Format: {image_format}",
        f"Dimensions: {width}x{height}",
    ]
    ocr_text = _run_tesseract(path)
    if ocr_text:
        lines.append("")
        lines.append("OCR text:")
        lines.append(ocr_text)
    else:
        lines.append("")
        lines.append("No readable OCR text was detected in the image.")
    return "\n".join(lines)


def _run_tesseract(path: Path) -> str:
    tesseract_path = shutil.which("tesseract")
    if not tesseract_path:
        return ""
    try:
        result = subprocess.run(
            [tesseract_path, str(path), "stdout", "--psm", "6"],
            capture_output=True,
            text=True,
            timeout=TESSERACT_TIMEOUT_SECONDS,
            check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return ""
    if result.returncode != 0:
        return ""
    return result.stdout.strip()


def _decode_tabular_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _stringify_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    if len(text) > 120:
        return text[:117].rstrip() + "..."
    return text
